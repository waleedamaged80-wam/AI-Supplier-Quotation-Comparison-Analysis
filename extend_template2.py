"""One-off: Submissions rows 5..204 -> 5..304 (10 suppliers x 30 lines)."""
import copy, sys
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.formatting.formatting import ConditionalFormattingList

src = sys.argv[1] if len(sys.argv) > 1 else 'P01113_template.xlsx'
wb = load_workbook(src)
sb = wb['Submissions']
for r in range(205, 305):
    sb.row_dimensions[r].height = sb.row_dimensions[204].height
    for c in range(1, 20):
        s, d = sb.cell(204, c), sb.cell(r, c)
        v = s.value
        d.value = Translator(v, origin=s.coordinate).translate_formula(d.coordinate) if isinstance(v, str) and v.startswith('=') else v
        d._style = copy.copy(s._style)
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and '$204' in c.value:
                c.value = c.value.replace('$204', '$304')
for dv in sb.data_validations.dataValidation:
    dv.sqref = type(dv.sqref)(str(dv.sqref).replace('204', '304'))
new = ConditionalFormattingList()
for cf in sb.conditional_formatting:
    for rule in cf.rules:
        new.add(str(cf.sqref).replace('204', '304'), rule)
sb.conditional_formatting = new
wb.save(src)
print('ok')
