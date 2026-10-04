"""RFQ / technical specification -> baseline tables (BOQ items + requirement lines), same format as the Baseline sheet.
The AI proposes; the user edits the tables before anything is used. Works with the API or Claude Code."""
import base64, io, json, os, re, shutil, subprocess, tempfile, time, zipfile
import extractor

SYSTEM = """You turn an RFQ and its technical specification (Arabic and/or English) into a compliance baseline for a procurement check.
Output two things:
1) BOQ items: each priced item in the RFQ (item number 1..3 at most, description, unit, quantity). If there are more than 3, keep the 3 largest and say so in the notes of the first line. Never invent items.
2) Requirement lines: one line per checkable requirement that a supplier quotation must meet, in the order: (a) one "Item N priced" line per BOQ item, (b) commercial lines (offer validity date, quotation date, prices ex-VAT, payment terms, warranty, origin, brands), (c) technical lines grouped by item.
Each line has:
- parameter: short, specific, includes the item or product when needed (e.g. "Windows: profile thickness (mm)").
- cls: "Mandatory" = failure makes the bid non-compliant (hard technical requirement, stated as must/minimum/shall); "Evaluated" = checked, deviation needs review; "Informational" = captured for comparison only (e.g. advance payment %).
- rule: "Min" (offer must be >= target, numbers/dates), "Max" (offer must be <= target), "Range" (offer within +/- tolerance of target), "Exact" (text must equal the target wording).
- target: a plain number (no units), a date as YYYY-MM-DD, or short text. For "Item N priced" lines the target is "Yes". For prices ex-VAT use "Ex-VAT". Leave empty only for Informational lines.
- tolerance: number or null. Use 0 for Min/Max/Range unless the RFQ allows a deviation.
- fail: "RED" for Mandatory lines, "YELLOW" for Evaluated and Informational lines.
- notes: where it comes from (e.g. "Spec s2: min 1.5 mm") and any assumption you made. Prefix assumptions with "ASSUMPTION:".
Rules: use only what the documents state; never invent requirements or quantities. If a date is relative, convert it only when the base date is given, otherwise leave target empty and say so in notes. Keep numbers in the unit named in the parameter. Do not copy prices from the RFQ.
Style example (existing baseline): "Windows: profile thickness (mm)" | Mandatory | Min | 1.5 | tol 0 | RED | "Spec s2: min 1.5 mm"; "Silicone brand" | Evaluated | Exact | "Wacker GP" | YELLOW | "Spec s5"; "Offer valid until" | Evaluated | Min | 2026-10-15 | YELLOW | "ASSUMPTION: must cover award window".
Return everything through the record_baseline tool."""


def build_tool():
    return {'name': 'record_baseline', 'description': 'Record the baseline extracted from the RFQ.',
            'input_schema': {'type': 'object', 'properties': {
                'rfq_id': {'type': ['string', 'null']}, 'title': {'type': ['string', 'null']},
                'boq': {'type': 'array', 'items': {'type': 'object', 'properties': {
                    'item': {'type': 'integer'}, 'description': {'type': 'string'}, 'unit': {'type': 'string'},
                    'qty': {'type': ['number', 'null']}}, 'required': ['item', 'description', 'unit', 'qty']}},
                'lines': {'type': 'array', 'items': {'type': 'object', 'properties': {
                    'parameter': {'type': 'string'},
                    'cls': {'type': 'string', 'enum': ['Mandatory', 'Evaluated', 'Informational']},
                    'rule': {'type': 'string', 'enum': ['Min', 'Max', 'Range', 'Exact']},
                    'target': {'type': ['string', 'number', 'null']},
                    'tolerance': {'type': ['number', 'null']},
                    'fail': {'type': 'string', 'enum': ['RED', 'YELLOW']},
                    'notes': {'type': 'string'}}, 'required': ['parameter', 'cls', 'rule', 'target', 'fail']}}},
                'required': ['boq', 'lines']}}


# ---------- reading the input ----------
def to_text(name, data):
    """Text of non-PDF inputs (docx, xlsx, csv, txt, md). Returns None for PDFs."""
    ext = os.path.splitext(name)[1].lower()
    if ext == '.pdf':
        return None
    if ext == '.docx':
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read('word/document.xml').decode('utf-8', 'ignore')
        xml = re.sub(r'</w:p>|</w:tr>', '\n', xml)
        xml = re.sub(r'</w:tc>', ' | ', xml)
        xml = re.sub(r'<w:tab/>', ' ', xml)
        text = re.sub(r'<[^>]+>', '', xml)
        for a, b in (('&amp;', '&'), ('&lt;', '<'), ('&gt;', '>'), ('&quot;', '"'), ('&apos;', "'")):
            text = text.replace(a, b)
        return text
    if ext in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True)
        out = []
        for ws in wb:
            out.append(f'## Sheet {ws.title}')
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    out.append(' | '.join('' if v is None else str(v) for v in row))
        return '\n'.join(out)
    return data.decode('utf-8', 'ignore')


def _prompt_head():
    return ('Build the compliance baseline from the attached RFQ / technical specification. '
            'Return the RFQ number and title if stated.')


# ---------- normalise ----------
def normalise(raw):
    """Tool output -> (rfq_id, title, boq_rows, line_rows) in the editor's column format."""
    boq = [dict(Item=int(b['item']), Description=b.get('description') or '', Unit=b.get('unit') or '',
                Qty=b.get('qty') or 0) for b in (raw.get('boq') or [])][:3]
    lines = []
    for i, l in enumerate(raw.get('lines') or [], 1):
        cls = l.get('cls') if l.get('cls') in ('Mandatory', 'Evaluated', 'Informational') else 'Evaluated'
        rule = l.get('rule') if l.get('rule') in ('Min', 'Max', 'Range', 'Exact') else 'Exact'
        t = l.get('target')
        if isinstance(t, float) and t.is_integer():
            t = int(t)
        tol = l.get('tolerance')
        if tol is None and rule in ('Min', 'Max', 'Range'):
            tol = 0
        lines.append({'Line': i, 'Parameter': str(l.get('parameter') or '').strip(), 'Class': cls, 'Rule': rule,
                      'Target': '' if t is None else str(t), 'Tolerance': tol,
                      'Fail result': l.get('fail') if l.get('fail') in ('RED', 'YELLOW') else
                      ('RED' if cls == 'Mandatory' else 'YELLOW'),
                      'Notes': (l.get('notes') or '')[:200]})
    return raw.get('rfq_id') or '', raw.get('title') or '', boq, lines[:30]


# ---------- API backend ----------
def extract_api(client, name, data, model=extractor.MODEL, retries=2):
    text = to_text(name, data)
    if text is None:
        content = [{'type': 'document', 'source': {'type': 'base64', 'media_type': 'application/pdf',
                                                   'data': base64.standard_b64encode(data).decode()}}]
    else:
        content = [{'type': 'text', 'text': f'<rfq_document name="{name}">\n{text[:150000]}\n</rfq_document>'}]
    content.append({'type': 'text', 'text': _prompt_head()})
    req = dict(model=model, max_tokens=8000, system=SYSTEM, tools=[build_tool()],
               tool_choice={'type': 'tool', 'name': 'record_baseline'},
               messages=[{'role': 'user', 'content': content}])
    for attempt in range(retries + 1):
        try:
            resp = client.messages.create(**req)
            for b in resp.content:
                if b.type == 'tool_use' and b.name == 'record_baseline':
                    return normalise(b.input)
            raise RuntimeError('no tool_use block in response')
        except Exception:
            if attempt == retries:
                raise
            time.sleep(2 * (attempt + 1))


# ---------- Claude Code backend ----------
def extract_cli(name, data, retries=1, timeout=900):
    exe = shutil.which('claude')
    if not exe:
        raise RuntimeError('Claude Code is not installed or not on PATH.')
    workdir = tempfile.mkdtemp(prefix='rfq_b_')
    text = to_text(name, data)
    fname = 'rfq.pdf' if text is None else 'rfq.txt'
    with open(os.path.join(workdir, fname), 'wb') as fh:
        fh.write(data if text is None else text.encode('utf-8'))
    prompt = (SYSTEM.replace('Return everything through the record_baseline tool.', 'Respond with the structured JSON only.')
              + f'\n\n{_prompt_head()}\nThe document is the file {fname} in the current folder. '
              'Use the Read tool to open it and read every page, including scanned pages.')
    cmd = [exe, '-p', prompt, '--output-format', 'json', '--json-schema', json.dumps(build_tool()['input_schema']),
           '--allowedTools', 'Read']
    for attempt in range(retries + 1):
        try:
            p = subprocess.run(cmd, cwd=workdir, env=extractor._cli_env(), capture_output=True, text=True,
                               encoding='utf-8', timeout=timeout)
            try:
                resp = json.loads(p.stdout)
            except json.JSONDecodeError:
                raise RuntimeError((p.stdout or p.stderr or 'empty reply')[:500])
            if p.returncode != 0 or resp.get('is_error'):
                raise RuntimeError(str(resp.get('result') or p.stderr)[:500])
            return normalise(resp.get('structured_output') or extractor._json_from_text(resp.get('result')))
        except Exception:
            if attempt == retries:
                raise
            time.sleep(3)


def extract_baseline(name, data, backend, api_key=None):
    """Returns (rfq_id, title, boq_rows, line_rows)."""
    if backend == 'api':
        import anthropic
        client = anthropic.Anthropic(api_key=api_key) if api_key else anthropic.Anthropic()
        return extract_api(client, name, data)
    return extract_cli(name, data)
