"""Python port of the workbook rules (Submissions J-M and the Compliance Matrix).
The sheet formulas stay the source of truth for users; this port lets the pipeline
and tests verify them. Keep both in sync."""
from datetime import date, datetime, time
from openpyxl.utils.datetime import to_excel

GREEN, YELLOW, RED, CLEARED = 'GREEN', 'YELLOW', 'RED', 'CLEARED'
SUB_FIRST, SUB_LAST = 5, 204          # Submissions rows
BASE_FIRST, BASE_LAST = 20, 49        # Baseline requirement rows
DEFAULT_THRESHOLD = 0.8


def fix(v):
    """openpyxl reads small numbers as datetimes when a cell carries a date format.
    Real dates are serials above 40000; anything below is a number again."""
    if isinstance(v, time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    if isinstance(v, (datetime, date)):
        n = to_excel(v)
        if n <= 40000:
            return int(n) if float(n).is_integer() else float(n)
    return v


def _cmp(x):
    if isinstance(x, datetime):
        x = x.date()
    return x.toordinal() if isinstance(x, date) else x


def _isnum(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def _text(x):
    if x is None:
        return ''
    if isinstance(x, float) and x.is_integer():
        x = int(x)
    return ' '.join(str(x).split()).upper()


def rule_result(line, offered):
    """Mirrors Submissions!J. line = dict(cls, rule, target, tol, fail)."""
    cls = line['cls']
    if cls == 'Informational':
        return GREEN
    if offered is None or offered == '':
        return RED if cls == 'Mandatory' else YELLOW
    fail = line.get('fail') or RED
    rule, tol = line['rule'], line.get('tol') or 0
    t, o = _cmp(line['target']), _cmp(offered)
    if rule == 'Exact':
        return GREEN if _text(o) == _text(t) else fail
    if not (_isnum(o) and _isnum(t)):
        return fail
    if rule == 'Min':
        return GREEN if o >= t else (YELLOW if o >= t - tol else fail)
    if rule == 'Max':
        return GREEN if o <= t else (YELLOW if o <= t + tol else fail)
    if rule == 'Range':
        return GREEN if abs(o - t) <= tol else fail
    return 'CHECK RULE'


def review_class(rule_res, declaration, confidence=None, threshold=DEFAULT_THRESHOLD):
    """Mirrors Submissions!K. Low AI confidence overrides GREEN to YELLOW."""
    if declaration == 'Non-Compliant':
        return RED
    if declaration in ('Alternative Offer', 'Partially Compliant') and rule_res == GREEN:
        return YELLOW
    if rule_res == GREEN and _isnum(confidence) and confidence < threshold:
        return YELLOW
    return rule_res


def final_status(review, cleared):
    """Mirrors Submissions!L."""
    return GREEN if review == GREEN else (CLEARED if cleared else review)


def flag(declaration, rule_res, confidence=None, threshold=DEFAULT_THRESHOLD):
    """Mirrors Submissions!M."""
    if not declaration:
        return 'Declaration missing'
    if declaration == 'Not Stated':
        return 'Not stated in quotation'
    if _isnum(confidence) and confidence < threshold:
        return 'Low AI confidence - verify'
    if declaration == 'Fully Compliant' and rule_res != GREEN:
        return 'Declared Fully - evidence differs'
    return ''


def load_baseline(wb):
    """Read header, BOQ and requirement lines from the Baseline sheet."""
    bs = wb['Baseline']
    header = {bs.cell(r, 1).value: bs.cell(r, 2).value for r in range(3, 9)}
    thr = bs['B9'].value if bs['A9'].value else None
    boq = []
    for r in range(11, 17):
        if bs.cell(r, 1).value is not None:
            boq.append(dict(item=bs.cell(r, 1).value, desc=bs.cell(r, 2).value,
                            unit=bs.cell(r, 3).value, qty=bs.cell(r, 4).value))
    lines = {}
    for r in range(BASE_FIRST, BASE_LAST + 1):
        n = bs.cell(r, 1).value
        if n is None:
            continue
        lines[int(n)] = dict(line=int(n), parameter=bs.cell(r, 2).value, cls=bs.cell(r, 3).value,
                             rule=bs.cell(r, 4).value, target=fix(bs.cell(r, 5).value),
                             tol=bs.cell(r, 6).value, fail=bs.cell(r, 7).value, notes=bs.cell(r, 8).value)
    return dict(header=header, boq=boq, lines=lines,
                threshold=thr if _isnum(thr) else DEFAULT_THRESHOLD)


def read_exceptions(wb):
    """Set of (supplier, rev, line) cleared by an Accepted exception with an approver."""
    ex, out = wb['Exceptions'], set()
    for r in range(5, 105):
        sup, rev, ln, dec, appr = (ex.cell(r, c).value for c in (2, 3, 4, 7, 8))
        if sup and dec == 'Accepted' and appr:
            out.add((sup, rev, ln))
    return out


def matrix(rows, baseline, cleared):
    """rows: list of dict(supplier, rev, line, offered, declaration, confidence).
    Returns {supplier: dict(rev, statuses{line: status}, red, yellow, cleared, overall, eligible)}."""
    thr = baseline['threshold']
    latest, out = {}, {}
    for r in rows:
        latest[r['supplier']] = max(latest.get(r['supplier'], 0), r['rev'])
    for sup, rev in latest.items():
        st = {}
        for r in rows:
            if r['supplier'] == sup and r['rev'] == rev and r['line'] in baseline['lines']:
                rr = rule_result(baseline['lines'][r['line']], r['offered'])
                rc = review_class(rr, r['declaration'], r.get('confidence'), thr)
                st[r['line']] = final_status(rc, (sup, rev, r['line']) in cleared)
        full = {n: st.get(n, 'NO DATA') for n in baseline['lines']}
        red = sum(v == RED for v in full.values())
        yel = sum(v == YELLOW for v in full.values())
        overall = ('INCOMPLETE' if 'NO DATA' in full.values()
                   else RED if red else YELLOW if yel else GREEN)
        out[sup] = dict(rev=rev, statuses=full, red=red, yellow=yel,
                        cleared=sum(v == CLEARED for v in full.values()),
                        overall=overall, eligible='YES' if overall == GREEN else 'NO')
    return out
