"""End-to-end checks. Run:  python test_pipeline.py
Needs LibreOffice only for the recalculation test (set RECALC=/path/to/recalc.py, else that test is skipped)."""
import json, os, subprocess, sys, tempfile, types
from datetime import date
from openpyxl import load_workbook
import rules, run, extractor

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, 'P01113_template.xlsx')
MANIFEST = os.path.join(HERE, 'manifest_p01113.json')
BASE = rules.load_baseline(load_workbook(TEMPLATE))


def test_rules():
    L = BASE['lines']
    assert rules.rule_result(L[12], 1.5) == 'GREEN' and rules.rule_result(L[12], 1.4) == 'RED'   # Min, mandatory
    assert rules.rule_result(L[13], 20) == 'RED' and rules.rule_result(L[13], None) == 'RED'     # missing mandatory
    assert rules.rule_result(L[16], None) == 'YELLOW'                                            # missing evaluated
    assert rules.rule_result(L[4], date(2026, 9, 30)) == 'YELLOW'                                # expired validity
    assert rules.rule_result(L[4], date(2026, 10, 21)) == 'GREEN'
    assert rules.rule_result(L[8], 'wacker gp') == 'GREEN'                                       # Exact, case-insensitive
    assert rules.rule_result(L[10], None) == 'GREEN'                                             # informational
    # AI confidence below threshold turns GREEN into YELLOW; declarations override
    assert rules.review_class('GREEN', 'Fully Compliant', 0.7, 0.8) == 'YELLOW'
    assert rules.review_class('GREEN', 'Fully Compliant', 0.9, 0.8) == 'GREEN'
    assert rules.review_class('GREEN', 'Non-Compliant') == 'RED'
    assert rules.final_status('RED', True) == 'CLEARED' and rules.final_status('GREEN', True) == 'GREEN'
    print('rules ok')


def test_extractor_parsing():
    """Mock the API: verifies request building, tool parsing and type coercion without a key."""
    raw = dict(supplier_name='Falcon', rates=[dict(item=1, rate=370), dict(item=2, rate=600), dict(item=3, rate=650)],
               lines=[dict(line=4, value='2026-09-30', declaration='Fully Compliant', confidence=0.9, source_page=1),
                      dict(line=12, value='1.5 mm', declaration='Fully Compliant', confidence=1.4),
                      dict(line=20, value=2.5, declaration='Bogus', confidence=0.8),
                      dict(line=8, value=None, declaration='Fully Compliant', confidence=0.9)])
    block = types.SimpleNamespace(type='tool_use', name='record_quotation', input=raw)
    client = types.SimpleNamespace(messages=types.SimpleNamespace(create=lambda **kw: types.SimpleNamespace(content=[block])))
    pdf = os.path.join(tempfile.gettempdir(), 'one_page.pdf')
    from pypdf import PdfWriter
    w = PdfWriter(); w.add_blank_page(200, 200); w.write(pdf)
    out = extractor.extract_quotation(client, BASE, 'Falcon', pdf, [1, 1])
    by = {l['line']: l for l in out['lines']}
    assert len(out['lines']) == len(BASE['lines'])                       # every baseline line present
    assert by[4]['value'] == date(2026, 9, 30)                           # date target -> date
    assert by[12]['value'] == 1.5 and by[12]['confidence'] == 1.0        # "1.5 mm" -> 1.5, confidence clipped
    assert by[20]['declaration'] == 'Not Stated'                         # invalid enum -> Not Stated
    assert by[8]['declaration'] == 'Not Stated'                          # null value -> Not Stated
    assert by[1]['declaration'] == 'Not Stated' and by[1]['confidence'] == 0.0   # line not returned
    req = extractor.build_request(BASE, open(pdf, 'rb').read(), 'Falcon')
    assert req['tool_choice']['name'] == 'record_quotation' and 'Wacker GP' in req['messages'][0]['content'][1]['text']
    print('extractor ok')


def test_claude_code_backend():
    """Runs the subscription backend against a fake `claude` executable: checks the command,
    that API keys are stripped from the environment, and that structured output is parsed."""
    import stat
    d = tempfile.mkdtemp()
    out = json.dumps({'is_error': False, 'structured_output': {
        'supplier_name': 'Falcon', 'rates': [{'item': 1, 'rate': 370}],
        'lines': [{'line': 12, 'value': 1.5, 'declaration': 'Fully Compliant', 'confidence': 0.9}]}})
    script = os.path.join(d, 'claude')
    open(script, 'w').write('#!/bin/sh\n[ -n "$ANTHROPIC_API_KEY" ] && echo \'{"is_error":true,"result":"api key leaked"}\' && exit 0\n'
                            '[ -f quotation.pdf ] || { echo \'{"is_error":true,"result":"no pdf"}\'; exit 0; }\n'
                            "echo '" + out + "'\n")
    os.chmod(script, os.stat(script).st_mode | stat.S_IEXEC)
    pdf = os.path.join(d, 'q.pdf')
    from pypdf import PdfWriter
    w = PdfWriter(); w.add_blank_page(200, 200); w.write(pdf)
    old_path, old_key = os.environ['PATH'], os.environ.get('ANTHROPIC_API_KEY')
    os.environ['PATH'], os.environ['ANTHROPIC_API_KEY'] = d + os.pathsep + old_path, 'sk-test'
    try:
        res = extractor.extract_with_claude_code(BASE, 'Falcon', pdf, [1, 1], retries=0)
    finally:
        os.environ['PATH'] = old_path
        os.environ.pop('ANTHROPIC_API_KEY', None) if old_key is None else os.environ.__setitem__('ANTHROPIC_API_KEY', old_key)
    by = {l['line']: l for l in res['lines']}
    assert by[12]['value'] == 1.5 and by[1]['declaration'] == 'Not Stated' and len(res['lines']) == len(BASE['lines'])
    cmd = extractor.build_cli_command(BASE, 'Falcon')
    assert cmd[1] == '-p' and '--json-schema' in cmd and cmd[cmd.index('--allowedTools') + 1] == 'Read'
    print('claude-code backend ok')


def test_load_and_recalc():
    recalc = os.environ.get('RECALC')
    tmp = tempfile.mkdtemp()
    out = os.path.join(tmp, 'result.xlsx')
    args = types.SimpleNamespace(workbook=TEMPLATE, manifest=MANIFEST, out=out, json=[os.path.join(HERE, 'fixtures', '*.json')])
    run.cmd_load(args)
    if not recalc:
        print('load ok (recalc test skipped: set RECALC)'); return
    subprocess.run([sys.executable, recalc, out, '120'], check=True, capture_output=True)
    assert run.cmd_check(types.SimpleNamespace(workbook=out)) == 0       # sheet formulas == python rules
    v = load_workbook(out, data_only=True)
    cm = v['Compliance Matrix']
    got = {cm.cell(4, c).value: (cm.cell(38, c).value, cm.cell(39, c).value) for c in range(5, 12)}
    assert got == {'Tilal Al-Eamar': (3, 11), 'Lamset Ebdaa': (5, 6), 'Falcon': (1, 6), 'Mishkal': (2, 6),
                   'Qimam': (0, 7), 'Al Shareef': (3, 11), 'Ultra Frame': (0, 6)}, got
    ev = v['Evaluation']
    assert ev['F9'].value == 182735 and ev['F13'].value == 219840 and ev['F14'].value == 'Incomplete'
    print('load + recalc + check ok')


if __name__ == '__main__':
    test_rules(); test_extractor_parsing(); test_claude_code_backend(); test_load_and_recalc()
