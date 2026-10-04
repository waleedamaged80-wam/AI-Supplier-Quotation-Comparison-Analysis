"""Engine behind the Streamlit app. No Streamlit imports, so it can be tested on its own.

Flow: baseline tables -> extract PDFs -> editable rows -> rules (rules.py) -> evaluation -> Excel export.
The Excel export reuses run.cmd_load, so the downloaded workbook has the same formulas as the desktop version.
"""
import math
import os
import tempfile
import types
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from pypdf import PdfReader

import extractor
import rules
import run

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, 'P01113_template.xlsx')
MAX_SUPPLIERS, MAX_ITEMS, MAX_LINES = 10, 3, 30
MAX_API_PAGES = 100
DECLARATIONS = extractor.DECLARATIONS
CLASSES = ['Mandatory', 'Evaluated', 'Informational']
RULES = ['Min', 'Max', 'Range', 'Exact']
FAILS = ['RED', 'YELLOW']
DECISIONS = ['Accepted', 'Rejected', 'Clarify']
ROW_COLS = ['Supplier', 'Line', 'Parameter', 'Offered', 'Declaration', 'Confidence', 'Reviewed', 'Source', 'Note']


# ---------- small helpers ----------
def clean(v):
    """Turn pandas NaN / NA into None."""
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    try:
        import pandas as pd
        if v is pd.NA or v is pd.NaT:
            return None
    except Exception:
        pass
    return v


def fmt(v):
    """Value -> text for the editor."""
    v = rules.fix(clean(v))
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def parse_target(s):
    """Editor text -> number, date or text."""
    s = ('' if s is None else str(s)).strip()
    if s == '':
        return None
    if s.lower() not in ('nan', 'inf', '-inf', 'infinity', '-infinity'):
        try:
            f = float(s.replace(',', ''))
            return int(f) if f.is_integer() else f
        except ValueError:
            pass
    try:
        return date.fromisoformat(s)
    except ValueError:
        return s


# ---------- baseline ----------
def baseline_tables(baseline):
    boq = [dict(Item=b['item'], Description=b['desc'], Unit=b['unit'], Qty=b['qty']) for b in baseline['boq']]
    lines = [{'Line': n, 'Parameter': l['parameter'], 'Class': l['cls'], 'Rule': l['rule'], 'Target': fmt(l['target']),
              'Tolerance': l.get('tol'), 'Fail result': l.get('fail') or 'RED', 'Notes': l.get('notes') or ''}
             for n, l in baseline['lines'].items()]
    return boq, lines


def make_baseline(rfq_id, title, boq_rows, line_rows, threshold=rules.DEFAULT_THRESHOLD):
    boq, lines = [], {}
    for r in boq_rows:
        item = clean(r.get('Item'))
        if item is None:
            continue
        boq.append(dict(item=int(item), desc=clean(r.get('Description')) or '', unit=clean(r.get('Unit')) or '',
                        qty=clean(r.get('Qty')) or 0))
    dup = 0
    for r in line_rows:
        n, p = clean(r.get('Line')), clean(r.get('Parameter'))
        if n is None or not p:
            continue
        if int(n) in lines:
            dup += 1
        lines[int(n)] = dict(line=int(n), parameter=str(p), cls=clean(r.get('Class')) or 'Evaluated',
                             rule=clean(r.get('Rule')) or 'Exact', target=parse_target(clean(r.get('Target'))),
                             tol=clean(r.get('Tolerance')), fail=clean(r.get('Fail result')) or 'RED',
                             notes=clean(r.get('Notes')) or '')
    bl = dict(header={'RFQ ID': rfq_id, 'Title': title}, boq=boq, lines=lines, threshold=threshold)
    bl['_dups'] = dup
    return bl


def validate(baseline):
    errs = []
    ids = sorted(b['item'] for b in baseline['boq'])
    if not ids:
        errs.append('Add at least one BOQ item.')
    elif len(ids) > MAX_ITEMS:
        errs.append(f'At most {MAX_ITEMS} BOQ items are supported.')
    elif ids != list(range(1, len(ids) + 1)):
        errs.append('Number the BOQ items 1, 2, 3 without gaps.')
    if not baseline['lines']:
        errs.append('Add at least one requirement line.')
    if len(baseline['lines']) > MAX_LINES:
        errs.append(f'At most {MAX_LINES} requirement lines are supported.')
    if baseline.get('_dups'):
        errs.append('Requirement line numbers must be unique.')
    for n, l in baseline['lines'].items():
        if l['rule'] in ('Min', 'Max', 'Range') and l['cls'] != 'Informational' and not isinstance(
                rules.fix(l['target']), (int, float, date, datetime)):
            errs.append(f'Line {n}: rule {l["rule"]} needs a number or a date (YYYY-MM-DD) as target.')
        if l['cls'] != 'Informational' and l['target'] is None:
            errs.append(f'Line {n}: add a target value.')
    return errs


# ---------- extraction ----------
def extract_all(files, baseline, backend, api_key=None, progress=None):
    """files: list of (supplier, filename, bytes). Returns (docs, errors) with docs in the same order."""
    tmp = tempfile.mkdtemp(prefix='rfq_')
    jobs, errors = [], []
    for i, (supplier, fname, data) in enumerate(files):
        path = os.path.join(tmp, f'{i:02d}.pdf')
        with open(path, 'wb') as fh:
            fh.write(data)
        try:
            n = len(PdfReader(path).pages)
        except Exception:
            errors.append((supplier, f'{fname} is not a readable PDF'))
            continue
        if backend == 'api' and n > MAX_API_PAGES:
            errors.append((supplier, f'{fname} has {n} pages; the API limit is {MAX_API_PAGES}. Split the PDF.'))
            continue
        jobs.append((i, supplier, path, n))

    client = None
    if backend == 'api' and jobs:
        import anthropic
        client = anthropic.Anthropic(api_key=api_key) if api_key else anthropic.Anthropic()

    def one(job):
        _, s, path, n = job
        if backend == 'api':
            return extractor.extract_quotation(client, baseline, s, path, [1, n])
        return extractor.extract_with_claude_code(baseline, s, path, [1, n])

    results, done = {}, 0
    workers = 3 if backend == 'api' else 1
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futs = {ex.submit(one, j): j for j in jobs}
        for f in as_completed(futs):
            j = futs[f]
            done += 1
            try:
                results[j[0]] = f.result()
            except Exception as e:                      # keep going with the other suppliers
                errors.append((j[1], str(e)[:300]))
            if progress:
                progress(done, len(jobs), j[1])
    docs = [results[j[0]] for j in jobs if j[0] in results]
    return docs, errors


# ---------- editable rows ----------
def make_rows(docs, baseline):
    rows = []
    for d in docs:
        by = {l['line']: l for l in d['lines']}
        for n, bl in baseline['lines'].items():
            l = by.get(n)
            if l is None:
                rows.append(dict(Supplier=d['supplier'], Line=n, Parameter=bl['parameter'], Offered='',
                                 Declaration='Not Stated', Confidence=0.0, Reviewed=False, Source='',
                                 Note='Not in extraction - run the comparison again'))
                continue
            src = f"p{l.get('source_page') or '?'}: {l.get('source_text', '')}" \
                if (l.get('source_text') or l.get('source_page')) else ''
            rows.append(dict(Supplier=d['supplier'], Line=n, Parameter=bl['parameter'], Offered=fmt(l['value']),
                             Declaration=l['declaration'], Confidence=float(l['confidence']), Reviewed=False,
                             Source=src, Note=l.get('note') or ''))
    return rows


def _offered(row, baseline):
    bl = baseline['lines'].get(int(row['Line']))
    text = clean(row.get('Offered'))
    if bl is None or text is None or str(text).strip() == '':
        return None
    return extractor._coerce(str(text).strip(), bl)


def rule_rows(rows, baseline):
    out = []
    for r in rows:
        n = int(r['Line'])
        if n not in baseline['lines']:
            continue
        conf = 1.0 if r.get('Reviewed') else float(clean(r.get('Confidence')) or 0)
        out.append(dict(supplier=r['Supplier'], rev=1, line=n, offered=_offered(r, baseline),
                        declaration=clean(r.get('Declaration')) or 'Not Stated', confidence=conf))
    return out


def compute(rows, baseline, accepted=()):
    """Returns (result with exceptions applied, result without). accepted = {(supplier, line)}."""
    rr = rule_rows(rows, baseline)
    cleared = {(s, 1, n) for s, n in accepted}
    return rules.matrix(rr, baseline, cleared), rules.matrix(rr, baseline, set())


def open_items(base_result, rows, baseline):
    out, by = [], {(r['Supplier'], int(r['Line'])): r for r in rows}
    for s, res in base_result.items():
        for n, st in res['statuses'].items():
            if st in ('RED', 'YELLOW'):
                r = by.get((s, n), {})
                out.append(dict(Supplier=s, Line=n, Parameter=baseline['lines'][n]['parameter'], Status=st,
                                Required=fmt(baseline['lines'][n]['target']), Offered=r.get('Offered', ''),
                                Note=r.get('Note', '')))
    return out


def accepted_set(decisions):
    """decisions: {(supplier, line): dict(Decision, Approver, Justification)}"""
    return {k for k, v in decisions.items() if v.get('Decision') == 'Accepted' and str(v.get('Approver') or '').strip()}


# ---------- evaluation ----------
def evaluate(suppliers, rates, adj, baseline, result):
    out = []
    for s in suppliers:
        rs = rates.get(s, {})
        vals = [clean(rs.get(b['item'])) for b in baseline['boq']]
        complete = all(v is not None for v in vals)
        quoted = sum(v * b['qty'] for v, b in zip(vals, baseline['boq'])) if complete else None
        a = clean(adj.get(s)) or 0
        indicative = quoted + a if complete else None
        eligible = result.get(s, {}).get('eligible') == 'YES'
        out.append(dict(Supplier=s, Eligible='YES' if eligible else 'NO', Quoted=quoted, Adjustments=a,
                        Indicative=indicative, Evaluated=indicative if (eligible and complete) else None))
    ev = [o['Evaluated'] for o in out if o['Evaluated'] is not None]
    ind = [o['Indicative'] for o in out if o['Indicative'] is not None]
    for o in out:
        o['Rank'] = 1 + sum(v < o['Evaluated'] for v in ev) if o['Evaluated'] is not None else None
        o['Indicative rank'] = 1 + sum(v < o['Indicative'] for v in ind) if o['Indicative'] is not None else None
    return out


# ---------- Excel export ----------
_SAMPLE_NOTES = {'SAMPLE EXERCISE', 'Data', 'Declaration', 'Check before use'}


def prepare_template(baseline, template=TEMPLATE):
    wb = load_workbook(template)
    bs = wb['Baseline']
    bs['B3'], bs['B4'] = baseline['header'].get('RFQ ID'), baseline['header'].get('Title')
    bs['B6'] = date.today()
    bs['B6'].number_format = run.DATE_FMT
    bs['B9'] = baseline['threshold']
    for r in range(11, 17):
        for c in range(1, 5):
            bs.cell(r, c).value = None
    for i, b in enumerate(baseline['boq'], 11):
        for c, v in enumerate((b['item'], b['desc'], b['unit'], b['qty']), 1):
            bs.cell(i, c).value = v
    for r in range(rules.BASE_FIRST, rules.BASE_LAST + 1):
        for c in range(1, 9):
            bs.cell(r, c).value = None
        bs.cell(r, 5).fill = PatternFill(fill_type=None)
    for i, (n, l) in enumerate(baseline['lines'].items(), rules.BASE_FIRST):
        for c, v in enumerate((n, l['parameter'], l['cls'], l['rule'], l['target'], l['tol'], l['fail'], l['notes']), 1):
            bs.cell(i, c).value = v
    sb = wb['Submissions']
    for r in range(rules.SUB_FIRST, rules.SUB_LAST + 1):
        for col in 'ABCGHIRS':
            sb[f'{col}{r}'].value = None
    cm = wb['Compliance Matrix']
    for c in run.SUP_COLS:
        cm[f'{c}4'].value = None
    ev = wb['Evaluation']
    for r in range(9, 19):
        for col in 'CDEGL':
            ev[f'{col}{r}'].value = None
    ex = wb['Exceptions']
    for r in range(5, 105):
        for col in 'ABCDFGHIJ':
            ex[f'{col}{r}'].value = None
    cl = wb['Clarifications']
    for r in range(4, 54):
        for c in range(1, 11):
            cl.cell(r, c).value = None
    ins = wb['Instructions']
    for r in range(1, 60):
        if ins.cell(r, 1).value in _SAMPLE_NOTES:
            ins.cell(r, 1).value = ins.cell(r, 2).value = None
    return wb


def build_workbook(baseline, docs, rows, rates, adj, decisions):
    """Returns the .xlsx bytes. rows are the edited editor rows; decisions {(supplier, line): {...}}."""
    tmp = tempfile.mkdtemp(prefix='rfq_x_')
    prepared = os.path.join(tmp, 'prepared.xlsx')
    prepare_template(baseline).save(prepared)

    import json
    by = {(r['Supplier'], int(r['Line'])): r for r in rows}
    suppliers = [d['supplier'] for d in docs]
    for i, d in enumerate(docs):
        orig = {l['line']: l for l in d['lines']}
        lines = []
        for n, bl in baseline['lines'].items():
            r, o = by.get((d['supplier'], n)), orig.get(n, {})
            if r is None:
                continue
            val = _offered(r, baseline)
            lines.append(dict(line=n, value=val, declaration=clean(r.get('Declaration')) or 'Not Stated',
                              confidence=1.0 if r.get('Reviewed') else float(clean(r.get('Confidence')) or 0),
                              source_page=o.get('source_page'), source_text=o.get('source_text', ''),
                              note=clean(r.get('Note')) or ''))
        rt = [dict(item=k, rate=v) for k, v in (rates.get(d['supplier']) or {}).items() if v is not None]
        with open(os.path.join(tmp, f'{i:02d}.json'), 'w', encoding='utf-8') as fh:
            json.dump(dict(supplier=d['supplier'], lines=lines, rates=rt), fh, ensure_ascii=False, default=str)

    out = os.path.join(tmp, 'result.xlsx')
    run.cmd_load(types.SimpleNamespace(workbook=prepared, json=[os.path.join(tmp, '*.json')], out=out, manifest=None))

    wb = load_workbook(out)
    ev = wb['Evaluation']
    for k, s in enumerate(suppliers):
        a = clean(adj.get(s)) or 0
        if a:
            ev[f'G{9 + k}'] = a
    ex = wb['Exceptions']
    n = 0
    for (s, ln), v in decisions.items():
        if not v.get('Decision'):
            continue
        r = 5 + n
        n += 1
        bl = baseline['lines'].get(ln, {})
        ex[f'A{r}'], ex[f'B{r}'], ex[f'C{r}'], ex[f'D{r}'] = f'EX-{n:03d}', s, 1, ln
        ex[f'F{r}'] = f"{bl.get('parameter', '')}: required {fmt(bl.get('target'))}, offered {v.get('Offered', '')}"
        ex[f'G{r}'], ex[f'H{r}'] = v.get('Decision'), v.get('Approver') or None
        ex[f'I{r}'] = date.today()
        ex[f'J{r}'] = v.get('Justification') or None
    wb.save(out)
    with open(out, 'rb') as fh:
        return fh.read()
