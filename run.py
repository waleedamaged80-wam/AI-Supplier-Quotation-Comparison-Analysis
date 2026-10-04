"""RFQ compliance pipeline.
  python run.py extract --workbook template.xlsx --manifest manifest.json --pdf binder.pdf --out extracted/
  python run.py load    --workbook template.xlsx --json extracted/*.json --out result.xlsx
  python run.py check   --workbook result.xlsx      (needs a recalculated file)
"""
import argparse, csv, glob, json, os, sys
from datetime import date, datetime
from openpyxl import load_workbook
from copy import copy
from openpyxl.styles import Alignment, Font, PatternFill
import rules

A = 'Arial'
BLUE, BLACK = Font(name=A, size=10, color='0000FF'), Font(name=A, size=10)
SUP_COLS = 'EFGHIJKL'                      # Compliance Matrix supplier columns


# ---------- workbook prep ----------
DATE_FMT = 'dd-mmm-yyyy'


def is_date(v):
    return isinstance(rules.fix(v), (date, datetime))


def normalise_formats(wb):
    """Plain numbers use General; only cells that hold dates carry a date format.
    (Mixed conditional formats confuse pandas/openpyxl, which then read numbers as dates.)"""
    bs, sb, cm = wb['Baseline'], wb['Submissions'], wb['Compliance Matrix']
    base = rules.load_baseline(wb)['lines']
    for n, l in base.items():
        r = rules.BASE_FIRST + [k for k in base].index(n)
        bs[f'E{r}'].value = l['target']
        bs[f'E{r}'].number_format = DATE_FMT if is_date(l['target']) else 'General'
        cm[f'D{7 + [k for k in base].index(n)}'].number_format = DATE_FMT if is_date(l['target']) else 'General'
    for r in range(rules.SUB_FIRST, rules.SUB_LAST + 1):
        for col in 'FG':
            sb[f'{col}{r}'].number_format = 'General'
        ln = sb.cell(r, 3).value
        if ln in base and is_date(base[ln]['target']):
            sb[f'F{r}'].number_format = DATE_FMT
            sb[f'G{r}'].number_format = DATE_FMT
        v = rules.fix(sb[f'G{r}'].value)
        if sb[f'A{r}'].value and isinstance(v, (date, datetime)):
            sb[f'G{r}'].value = v          # restore date; numbers were re-typed by fix()
        elif sb[f'A{r}'].value and sb[f'G{r}'].value is not None:
            sb[f'G{r}'].value = v


def fix_compat(wb):
    """MAXIFS needs Excel 2019/365. Replace with SUMPRODUCT(MAX()), which works in every version."""
    cm = wb['Compliance Matrix']
    for c in SUP_COLS:
        f = cm[f'{c}5'].value
        if isinstance(f, str) and 'MAXIFS' in f:
            cm[f'{c}5'] = (f'=IF({c}4="","",IF(COUNTIF(Submissions!$A$5:$A$204,{c}4)=0,"",'
                           f'SUMPRODUCT(MAX((Submissions!$A$5:$A$204={c}4)*Submissions!$B$5:$B$204))))')


def ensure_ai_columns(wb):
    """Adds AI confidence + source columns and the low-confidence rule (idempotent)."""
    normalise_formats(wb)
    fix_compat(wb)
    bs, sb = wb['Baseline'], wb['Submissions']
    if not bs['A9'].value:
        bs['A9'], bs['B9'] = 'AI confidence threshold', rules.DEFAULT_THRESHOLD
        bs['A9'].font = Font(name=A, size=10, bold=True)
        bs['B9'].font, bs['B9'].number_format = BLUE, '0%'
        bs['B9'].fill = PatternFill('solid', fgColor='FFFF00')
        bs['B9'].alignment = Alignment(horizontal='left')
        bs['C9'], bs['C9'].font = 'Lines below this need human review', Font(name=A, size=9, italic=True, color='7F7F7F')
    for col, text in (('R', 'AI conf.'), ('S', 'Source (page: quote)')):
        c = sb[f'{col}4']
        c.value, c.font, c.fill, c.border = text, copy(sb['A4'].font), copy(sb['A4'].fill), copy(sb['A4'].border)
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in range(rules.SUB_FIRST, rules.SUB_LAST + 1):
        sb[f'K{r}'] = (f'=IF($A{r}="","",IF($H{r}="Non-Compliant","RED",IF(AND(OR($H{r}="Alternative Offer",$H{r}="Partially Compliant"),'
                       f'$J{r}="GREEN"),"YELLOW",IF(AND($J{r}="GREEN",ISNUMBER($R{r}),$R{r}<Baseline!$B$9),"YELLOW",$J{r}))))')
        sb[f'M{r}'] = (f'=IF($A{r}="","",IF($H{r}="","Declaration missing",IF($H{r}="Not Stated","Not stated in quotation",'
                       f'IF(AND(ISNUMBER($R{r}),$R{r}<Baseline!$B$9),"Low AI confidence - verify",'
                       f'IF(AND($H{r}="Fully Compliant",$J{r}<>"GREEN"),"Declared Fully - evidence differs","")))))')
        sb[f'R{r}'].font, sb[f'R{r}'].number_format = BLUE, '0%'
        sb[f'S{r}'].font = Font(name=A, size=9, color='7F7F7F')
    sb.column_dimensions['R'].width, sb.column_dimensions['S'].width = 9, 60


# ---------- extract ----------
def cmd_extract(a):
    import extractor, shutil
    wb = load_workbook(a.workbook)
    baseline = rules.load_baseline(wb)
    manifest = json.load(open(a.manifest, encoding='utf-8'))
    os.makedirs(a.out, exist_ok=True)
    client = None
    if a.backend == 'api':
        if not a.dry_run:
            import anthropic
            client = anthropic.Anthropic()            # reads ANTHROPIC_API_KEY, billed per use
    elif not a.dry_run and not shutil.which('claude'):
        sys.exit('Claude Code not found. Install it, run "claude" once to log in, then retry.')
    elif os.environ.get('ANTHROPIC_API_KEY'):
        print('note: ANTHROPIC_API_KEY is set; it is ignored here so your subscription is used.')
    for q in manifest['quotations']:
        path = a.pdf or q['file']
        if a.dry_run:
            size = len(extractor.slice_pdf(path, *q['pages'])) / 1e6
            print(f"[dry-run] {q['supplier']:<16} pages {q['pages']}  pdf {size:.2f} MB  backend {a.backend}")
            continue
        print(f"extracting {q['supplier']} ...", flush=True)
        if a.backend == 'api':
            data = extractor.extract_quotation(client, baseline, q['supplier'], path, q['pages'])
        else:
            data = extractor.extract_with_claude_code(baseline, q['supplier'], path, q['pages'])
        slug = q['supplier'].lower().replace(' ', '_').replace('-', '_')
        with open(os.path.join(a.out, f'{slug}.json'), 'w', encoding='utf-8') as fh:
            json.dump(data, fh, indent=1, ensure_ascii=False, default=str)


# ---------- load ----------
def _val(v):
    if isinstance(v, str) and len(v) == 10 and v[4] == '-' and v[7] == '-':
        try:
            return date.fromisoformat(v)
        except ValueError:
            pass
    return v


def cmd_load(a):
    wb = load_workbook(a.workbook)
    ensure_ai_columns(wb)
    sb, cm, ev = wb['Submissions'], wb['Compliance Matrix'], wb['Evaluation']
    baseline = rules.load_baseline(wb)
    thr = baseline['threshold']
    nxt = rules.SUB_FIRST
    revs = {}
    for r in range(rules.SUB_FIRST, rules.SUB_LAST + 1):
        s = sb.cell(r, 1).value
        if s:
            nxt = r + 1
            revs[s] = max(revs.get(s, 0), sb.cell(r, 2).value or 0)
    review = []
    files = [f for pat in a.json for f in sorted(glob.glob(pat))]
    docs = [json.load(open(f, encoding='utf-8')) for f in files]
    if a.manifest:
        order = [q['supplier'] for q in json.load(open(a.manifest, encoding='utf-8'))['quotations']]
        docs.sort(key=lambda d: order.index(d['supplier']) if d['supplier'] in order else 99)
    for d in docs:
        sup = d['supplier']
        rev = revs.get(sup, 0) + 1                # append-only: a returning supplier gets the next revision
        revs[sup] = rev
        if nxt + len(d['lines']) - 1 > rules.SUB_LAST:
            sys.exit('Submissions sheet is full')
        for ln in d['lines']:
            r = nxt; nxt += 1
            v = _val(ln['value'])
            sb[f'A{r}'], sb[f'B{r}'], sb[f'C{r}'] = sup, rev, ln['line']
            sb[f'F{r}'].number_format = DATE_FMT if is_date(baseline['lines'][ln['line']]['target']) else 'General'
            sb[f'H{r}'] = ln['declaration']
            if v is not None:
                sb[f'G{r}'] = v
                if isinstance(v, date):
                    sb[f'G{r}'].number_format = DATE_FMT
                sb[f'R{r}'] = ln['confidence']
                if ln['confidence'] < thr:
                    review.append([sup, rev, ln['line'], baseline['lines'][ln['line']]['parameter'], v,
                                   ln['confidence'], ln.get('source_page'), ln.get('source_text', '')])
            if ln.get('note'):
                sb[f'I{r}'] = ln['note']
            if ln.get('source_text') or ln.get('source_page'):
                sb[f'S{r}'] = f"p{ln.get('source_page') or '?'}: {ln.get('source_text', '')}"[:250]
            for col in 'ABCGHI':
                sb[f'{col}{r}'].font = BLUE
        # supplier slot on the matrix + rates on the evaluation sheet
        slot = next((c for c in SUP_COLS if cm[f'{c}4'].value == sup), None)
        if slot is None:
            slot = next((c for c in SUP_COLS if not cm[f'{c}4'].value), None)
            if slot is None:
                sys.exit('Compliance Matrix supports 8 suppliers')
            cm[f'{slot}4'] = sup
            cm[f'{slot}4'].font = Font(name=A, size=10, color='0000FF', bold=True)
        row = 9 + SUP_COLS.index(slot)
        for x in d.get('rates', []):
            if 1 <= x['item'] <= 3 and x.get('rate') is not None:
                ev[f"{'CDE'[x['item'] - 1]}{row}"] = x['rate']
        print(f"loaded {sup} rev {rev}: {len(d['lines'])} lines")
    wb.save(a.out)
    if review:
        qpath = os.path.splitext(a.out)[0] + '_review_queue.csv'
        with open(qpath, 'w', newline='', encoding='utf-8-sig') as fh:
            w = csv.writer(fh)
            w.writerow(['supplier', 'rev', 'line', 'parameter', 'value', 'confidence', 'page', 'source_text'])
            w.writerows(review)
        print(f'{len(review)} low-confidence values -> {qpath}')
    print(f'saved {a.out}. Open in Excel (or recalc with LibreOffice) to refresh formulas.')


# ---------- check ----------
def cmd_check(a):
    """Compare the workbook's computed statuses (Submissions!L) with the Python rules."""
    wb = load_workbook(a.workbook)                       # formulas
    wv = load_workbook(a.workbook, data_only=True)       # cached values
    baseline, cleared = rules.load_baseline(wb), rules.read_exceptions(wb)
    sv = wv['Submissions']
    bad = n = 0
    for r in range(rules.SUB_FIRST, rules.SUB_LAST + 1):
        if not sv.cell(r, 1).value:
            continue
        sup, rev, ln, off, dec, conf = (sv.cell(r, c).value for c in (1, 2, 3, 7, 8, 18))
        off = rules.fix(off)
        rr = rules.rule_result(baseline['lines'][ln], off)
        rc = rules.review_class(rr, dec, conf, baseline['threshold'])
        py = rules.final_status(rc, (sup, rev, ln) in cleared)
        n += 1
        if py != sv.cell(r, 12).value:
            bad += 1
            print(f'MISMATCH row {r}: {sup} line {ln}: sheet={sv.cell(r, 12).value} python={py}')
    print(f'{n} rows checked, {bad} mismatches')
    return bad


def main():
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest='cmd', required=True)
    e = sub.add_parser('extract')
    e.add_argument('--workbook', required=True); e.add_argument('--manifest', required=True)
    e.add_argument('--pdf'); e.add_argument('--out', default='extracted'); e.add_argument('--dry-run', action='store_true')
    e.add_argument('--backend', choices=['claude-code', 'api'], default='claude-code',
                   help='claude-code = your Claude subscription (default); api = pay-per-use API credits')
    l = sub.add_parser('load')
    l.add_argument('--workbook', required=True); l.add_argument('--json', nargs='+', required=True)
    l.add_argument('--out', required=True); l.add_argument('--manifest', help='load suppliers in manifest order')
    c = sub.add_parser('check'); c.add_argument('--workbook', required=True)
    a = p.parse_args()
    {'extract': cmd_extract, 'load': cmd_load, 'check': cmd_check}[a.cmd](a)


if __name__ == '__main__':
    main()
